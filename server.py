import io
import math
import os
import socket
import tempfile
import time
import base64
import json
import hmac
import hashlib
from datetime import datetime
from zoneinfo import ZoneInfo
from typing import Optional

try:
    BOGOTA_TZ = ZoneInfo("America/Bogota")
except Exception:
    from datetime import timezone, timedelta
    BOGOTA_TZ = timezone(timedelta(hours=-5))

def ahora_bogota():
    return datetime.now(BOGOTA_TZ)

import pandas as pd
from fastapi import FastAPI, Request, HTTPException, Response, UploadFile, File, Depends
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

import database
import pdf_generator

app = FastAPI(title="TireAnalyze Pro - Sistema Integral Web & Móvil")

@app.on_event("startup")
def startup_migracion_produccion():
    try:
        database.migrar_pepsico_a_tellantas()
    except Exception as e:
        print(f"[STARTUP] Error ejecutando migración: {e}")

# Montar directorio de recursos estáticos (imágenes HD, CSS, logos)
os.makedirs("static", exist_ok=True)
app.mount("/static", StaticFiles(directory="static"), name="static")

# ==========================================
# MODELOS PYDANTIC
# ==========================================

class LoginItem(BaseModel):
    username: str
    password: str

class UsuarioItem(BaseModel):
    username: str
    password: str
    nombre_completo: str
    rol: str = "inspector"
    cliente_asignado: str = ""
    vistas_permitidas: Optional[str] = "patio,diagnostico,kpis"

class EditarUsuarioItem(BaseModel):
    nombre_completo: Optional[str] = None
    password: Optional[str] = None
    rol: Optional[str] = None
    cliente_asignado: Optional[str] = None
    activo: Optional[bool] = None
    vistas_permitidas: Optional[str] = None

class ResetPasswordItem(BaseModel):
    nueva_password: str

class ConfigLicenciaClienteItem(BaseModel):
    cliente: str
    tipo_licencia: str = "PRO"  # "LIGHT" o "PRO"
    plantilla_kpis_json: str = "{}"


class InspeccionItem(BaseModel):
    fecha: str = ""
    cliente: str = ""
    subcliente: str = ""
    tipo_operacion: str = ""
    placa: str
    posicion: int
    kilometraje: float = 0.0
    tipologia: str = "Cabezote"
    marca: str = ""
    diseno: str = ""
    dimension: str = "295/80R22.5"
    vida: str = "Original"
    banda_reen: str = ""
    prof_ext: float = 0.0
    prof_cent: float = 0.0
    prof_int: float = 0.0
    psi: float = 100.0
    novedades: str = ""
    precio: float = 0.0
    prof_inicial: float = 0.0
    km_montaje: float = 0.0
    numero_interno: str = ""
    usuario_registro: str = ""
    foto_novedad: str = ""
    tipo_licencia: Optional[str] = None

class NuevaOperacionItem(BaseModel):
    nombre: str

class EditarOperacionItem(BaseModel):
    nombre_antiguo: str
    nombre_nuevo: str

class NuevaTipologiaItem(BaseModel):
    nombre: str
    descripcion: str = ""
    cant_llantas: Optional[int] = None
    ejes_dir: int = 1
    tipo_dir: str = "sencilla"
    ejes_trac: int = 2
    tipo_trac: str = "dual"
    ejes_libres: int = 0
    tipo_libres: str = "dual"
    cant_repuestos: int = 0
    esquema_custom: Optional[list] = None

class EditarTipologiaItem(BaseModel):
    nombre_antiguo: str
    nombre_nuevo: str
    descripcion: Optional[str] = ""
    cant_llantas: Optional[int] = 10
    ejes_dir: Optional[int] = None
    tipo_dir: Optional[str] = "sencilla"
    ejes_trac: Optional[int] = None
    tipo_trac: Optional[str] = "dual"
    ejes_libres: Optional[int] = None
    tipo_libres: Optional[str] = "dual"
    cant_repuestos: Optional[int] = None
    esquema_custom: Optional[list] = None

class ModeloLlantaItem(BaseModel):
    tipo_llanta: str = "Nueva"
    marca: str
    diseno: str
    dimension: str
    prof_inicial: float
    descripcion: str = ""

class EditarModeloLlantaItem(BaseModel):
    id: int
    tipo_llanta: str = "Nueva"
    marca: str
    diseno: str
    dimension: str
    prof_inicial: float
    descripcion: Optional[str] = ""
    actualizar_bd: Optional[bool] = True

class LimiteClienteItem(BaseModel):
    cliente: str
    subcliente: str = "TODOS"
    eje: str = "TODOS"
    limite_retiro_mm: float = 3.0
    observaciones: str = ""

class EditarLimiteClienteItem(BaseModel):
    id: int
    cliente: str
    subcliente: Optional[str] = "TODOS"
    eje: Optional[str] = "TODOS"
    limite_retiro_mm: float = 3.0
    observaciones: Optional[str] = ""

class MarcaClienteItem(BaseModel):
    cliente: str
    marca: str

class ClienteDirectorioItem(BaseModel):
    cliente: str
    tipo_licencia: Optional[str] = "PRO"
    limite_retiro_mm: Optional[float] = 3.0
    subclientes: Optional[str] = ""
    observaciones: Optional[str] = ""

# ==========================================
# SEGURIDAD Y TOKENS STATELESS (HMAC-SHA256)
# ==========================================

SECRET_KEY = os.environ.get("TIREANALYZE_SECRET_KEY", os.environ.get("TIRECHECK_SECRET_KEY", "tireanalyze-pro-super-secret-key-2026-production"))
TOKEN_EXPIRATION_SECONDS = 43200  # 12 horas

def crear_token_acceso(username: str, rol: str) -> str:
    payload = {
        "sub": username,
        "rol": rol,
        "exp": int(time.time()) + TOKEN_EXPIRATION_SECONDS
    }
    payload_json = json.dumps(payload, separators=(',', ':'))
    payload_b64 = base64.urlsafe_b64encode(payload_json.encode('utf-8')).decode('utf-8').rstrip('=')
    signature = hmac.new(SECRET_KEY.encode('utf-8'), payload_b64.encode('utf-8'), hashlib.sha256).hexdigest()
    return f"{payload_b64}.{signature}"

def verificar_token(token: str) -> Optional[dict]:
    if not token or not isinstance(token, str):
        return None
    parts = token.strip().split('.')
    if len(parts) != 2:
        return None
    payload_b64, signature = parts
    expected_signature = hmac.new(SECRET_KEY.encode('utf-8'), payload_b64.encode('utf-8'), hashlib.sha256).hexdigest()
    if not hmac.compare_digest(signature, expected_signature):
        return None
    try:
        padding = '=' * ((4 - len(payload_b64) % 4) % 4)
        payload_bytes = base64.urlsafe_b64decode(payload_b64 + padding)
        payload = json.loads(payload_bytes.decode('utf-8'))
        if payload.get("exp", 0) < int(time.time()):
            return None
        return payload
    except Exception:
        return None

def obtener_usuario_actual(request: Request) -> dict:
    token = None
    auth_header = request.headers.get("Authorization")
    if auth_header and auth_header.startswith("Bearer "):
        token = auth_header[7:].strip()
    elif "tireanalyze_token" in request.cookies:
        token = request.cookies.get("tireanalyze_token")
    elif "tirecheck_token" in request.cookies:
        token = request.cookies.get("tirecheck_token")
    elif "token" in request.query_params:
        token = request.query_params.get("token")
    
    if not token:
        raise HTTPException(status_code=401, detail="Sesión no válida o expirada")
    
    payload = verificar_token(token)
    if not payload or not payload.get("sub"):
        raise HTTPException(status_code=401, detail="Sesión no válida o expirada")
    
    usr = database.obtener_usuario_por_username(payload["sub"])
    if not usr or not usr.get("activo"):
        raise HTTPException(status_code=401, detail="Usuario no encontrado o inactivo")
    
    return usr

def es_rol_admin(rol: str, user: str = "") -> bool:
    r = str(rol or "").strip().lower()
    u = str(user or "").strip().lower()
    return r in ("admin", "administrador") or u == "admin"

def es_rol_jefatura(rol: str) -> bool:
    r = str(rol or "").strip().lower()
    return r in ("jefatura", "jefe", "jefe_llantas")

def es_rol_gerencia(rol: str) -> bool:
    r = str(rol or "").strip().lower()
    return r in ("gerencia", "gerente")

def requerir_admin(usuario_actual: dict = Depends(obtener_usuario_actual)) -> dict:
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    user = str(usuario_actual.get("username", "")).strip().lower()
    if not es_rol_admin(rol, user):
        raise HTTPException(status_code=403, detail="Se requieren permisos de administrador")
    return usuario_actual

def requerir_admin_o_jefatura(usuario_actual: dict = Depends(obtener_usuario_actual)) -> dict:
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    user = str(usuario_actual.get("username", "")).strip().lower()
    if not (es_rol_admin(rol, user) or es_rol_jefatura(rol)):
        raise HTTPException(status_code=403, detail="Se requieren permisos de jefatura o administrador")
    return usuario_actual

def resolver_filtro_usuario_inspecciones(usuario_actual: dict, usuario_param: str = "") -> Optional[str]:
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    user = str(usuario_actual.get("username", "")).strip().lower()
    if es_rol_admin(rol, user):
        return usuario_param if usuario_param else None
    if es_rol_jefatura(rol) or es_rol_gerencia(rol):
        if usuario_param:
            return usuario_param
        elif usuario_actual.get("cliente_asignado"):
            return usuario_actual["username"]
        else:
            return None
    return usuario_actual["username"]


# ==========================================
# RUTAS PRINCIPALES Y SERVIDOR
# ==========================================

@app.get("/", response_class=HTMLResponse)
def index(response: Response):
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    template_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates", "index.html")
    with open(template_path, "r", encoding="utf-8") as f:
        return f.read()

@app.get("/preview-chasis", response_class=HTMLResponse)
def preview_chasis():
    template_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates", "preview_chasis.html")
    with open(template_path, "r", encoding="utf-8") as f:
        return f.read()

@app.get("/api/servidor/info")
def api_info_servidor():
    return {
        "app": "TireAnalyze Pro",
        "status": "online",
        "version": "2.0"
    }

def sanitizar_nativos(obj):
    if isinstance(obj, list):
        return [sanitizar_nativos(i) for i in obj]
    elif isinstance(obj, dict):
        return {k: sanitizar_nativos(v) for k, v in obj.items()}
    elif pd.isna(obj):
        return None
    elif hasattr(obj, "item"):
        val = obj.item()
        return None if (isinstance(val, float) and math.isnan(val)) else val
    elif isinstance(obj, float) and math.isnan(obj):
        return None
    return obj


# ==========================================
# RUTAS DE INSPECCIONES (CRUD & EXCEL)
# ==========================================

@app.get("/api/inspecciones")
def api_obtener_inspecciones(usuario: str = "", usuario_actual: dict = Depends(obtener_usuario_actual)):
    filtro_u = resolver_filtro_usuario_inspecciones(usuario_actual, usuario)
    df = database.obtener_todas_inspecciones(usuario=filtro_u)
    return sanitizar_nativos(df.to_dict(orient="records"))


@app.get("/api/inspecciones/plantilla-excel")
def api_descargar_plantilla_excel(usuario_actual: dict = Depends(obtener_usuario_actual)):
    _p_data = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "plantillas", "Plantilla_Cargue_Masivo_Inspecciones.xlsx")
    _p_root = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Plantilla_Cargue_Masivo_Inspecciones.xlsx")
    plantilla_path = _p_data if os.path.exists(_p_data) else _p_root
    if os.path.exists(plantilla_path):
        with open(plantilla_path, "rb") as f:
            data = f.read()
    else:
        columnas = [
            'fecha', 'cliente', 'subcliente', 'tipo_operacion', 'placa', 'numero_interno',
            'posicion', 'kilometraje', 'tipologia', 'marca', 'diseno', 'dimension',
            'vida', 'banda_reen', 'prof_ext', 'prof_cent', 'prof_int', 'psi',
            'novedades', 'precio', 'prof_inicial', 'km_montaje'
        ]
        datos_guia = [{
            'fecha': ahora_bogota().strftime('%Y-%m-%d'),
            'cliente': 'Axxom', 'subcliente': 'Principal', 'tipo_operacion': 'Carga Seca',
            'placa': 'ABC123', 'numero_interno': 'LL-001', 'posicion': 1, 'kilometraje': 120500,
            'tipologia': 'Cabezote (6x4)', 'marca': 'Michelin', 'diseno': 'X Multi Z',
            'dimension': '295/80R22.5', 'vida': 'Original', 'banda_reen': '',
            'prof_ext': 14.5, 'prof_cent': 14.2, 'prof_int': 14.4, 'psi': 110,
            'novedades': 'Buen estado', 'precio': 1650000, 'prof_inicial': 16.5, 'km_montaje': 80000
        }]
        df = pd.DataFrame(datos_guia, columns=columnas)
        buf = io.BytesIO()
        df.to_excel(buf, index=False, sheet_name='Plantilla_Inspecciones')
        data = buf.getvalue()

    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Plantilla_Cargue_Masivo_Inspecciones.xlsx"}
    )

@app.get("/api/inspecciones/exportar-excel")
def api_exportar_inspecciones_excel(usuario: str = "", usuario_actual: dict = Depends(obtener_usuario_actual)):
    filtro_u = resolver_filtro_usuario_inspecciones(usuario_actual, usuario)
    df = database.obtener_todas_inspecciones(usuario=filtro_u)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Inspecciones", index=False)
    f_str = ahora_bogota().strftime("%Y%m%d_%H%M")
    filename = f"Base_Inspecciones_{f_str}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@app.post("/api/inspecciones/importar-excel")
async def api_importar_excel(file: UploadFile = File(...), usuario_actual: dict = Depends(obtener_usuario_actual)):
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    if es_rol_gerencia(rol):
        raise HTTPException(status_code=403, detail="El rol Gerencia tiene perfil ejecutivo de consulta y no puede importar inspecciones")
    filename = file.filename or ""
    if not filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Formato no soportado. El archivo debe ser Excel (.xlsx o .xls)")

    content = await file.read()
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        tmp.write(content)
        tmp_path = tmp.name

    try:
        cant = database.cargar_excel_a_bd(tmp_path)
        return {"status": "ok", "registros_importados": cant, "message": f"Se importaron exitosamente {cant} registros a la base de datos."}
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error al procesar el archivo Excel: {str(e)}")
    finally:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except Exception:
                pass

@app.get("/api/inspecciones/ultima-marca-diseno")
def api_inspecciones_ultima_marca_diseno(placa: str, usuario_actual: dict = Depends(obtener_usuario_actual)):
    res = database.obtener_ultima_marca_diseno_vehiculo(placa)
    return sanitizar_nativos(res)

@app.get("/api/inspecciones/{inspeccion_id}")
def api_obtener_inspeccion_por_id(inspeccion_id: int, usuario_actual: dict = Depends(obtener_usuario_actual)):
    with database.get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT id, fecha, cliente, subcliente, tipo_operacion, placa, posicion, kilometraje,
                   tipologia, marca, diseno, dimension, vida, banda_reen, prof_ext, prof_cent, prof_int,
                   psi, novedades, precio, prof_inicial, km_montaje, numero_interno, usuario_registro, foto_novedad
            FROM inspecciones WHERE id = ?
        """, (inspeccion_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Inspección no encontrada")
        
        # Validar permisos de clientes asignados
        rol = str(usuario_actual.get("rol", "")).strip().lower()
        if not es_rol_admin(rol, usuario_actual.get("username")):
            clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
            if clis and row[2] and not any(c.lower() == str(row[2]).strip().lower() for c in clis):
                raise HTTPException(status_code=403, detail="No tiene permisos para ver inspecciones de este cliente")
        
        cols = ['id', 'fecha', 'cliente', 'subcliente', 'tipo_operacion', 'placa', 'posicion', 'kilometraje',
                'tipologia', 'marca', 'diseno', 'dimension', 'vida', 'banda_reen', 'prof_ext', 'prof_cent', 'prof_int',
                'psi', 'novedades', 'precio', 'prof_inicial', 'km_montaje', 'numero_interno', 'usuario_registro', 'foto_novedad']
        return dict(zip(cols, row))

@app.post("/api/inspecciones")
def api_guardar_inspeccion(item: InspeccionItem, usuario_actual: dict = Depends(obtener_usuario_actual)):
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    if es_rol_gerencia(rol):
        raise HTTPException(status_code=403, detail="El rol Gerencia tiene perfil ejecutivo de consulta y no puede registrar inspecciones")
    datos = item.model_dump()
    if not datos.get("fecha"):
        datos["fecha"] = ahora_bogota().strftime("%Y-%m-%d")

    # Aislamiento por rol y asignación de clientes
    if not es_rol_admin(rol, usuario_actual.get("username")):
        datos["usuario_registro"] = usuario_actual["username"]
        clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
        if clis:
            # Si solo tiene un cliente asignado, predeterminarlo automáticamente
            if not datos.get("cliente") and len(clis) == 1:
                datos["cliente"] = clis[0]
            elif not any(c.lower() == str(datos.get("cliente", "")).strip().lower() for c in clis):
                raise HTTPException(status_code=403, detail="No tiene permisos para registrar inspecciones de este cliente")

    # Validación de Candados de Seguridad
    val = database.validar_inspeccion_candados(datos)
    if not val.get('valido', True):
        raise HTTPException(status_code=400, detail=val.get('mensaje', 'Violación de regla de seguridad'))

    reg_id = database.insertar_inspeccion(datos)
    return {"status": "ok", "id": reg_id, "message": "Inspección guardada correctamente"}

@app.put("/api/inspecciones/{inspeccion_id}")
def api_actualizar_inspeccion(inspeccion_id: int, item: InspeccionItem, usuario_actual: dict = Depends(obtener_usuario_actual)):
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    if es_rol_gerencia(rol):
        raise HTTPException(status_code=403, detail="El rol Gerencia tiene perfil ejecutivo de consulta y no puede modificar inspecciones")
    datos = item.model_dump()
    
    # Validar permisos de asignación de cliente
    if not es_rol_admin(rol, usuario_actual.get("username")):
        with database.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT cliente, usuario_registro FROM inspecciones WHERE id = ?", (inspeccion_id,))
            row_exist = cursor.fetchone()
            if not row_exist:
                raise HTTPException(status_code=404, detail="No se encontró el registro para actualizar")
            clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
            if clis and row_exist[0] and not any(c.lower() == str(row_exist[0]).strip().lower() for c in clis):
                raise HTTPException(status_code=403, detail="No tiene permisos para modificar inspecciones de este cliente")
        datos["usuario_registro"] = usuario_actual["username"]
        if clis:
            if not datos.get("cliente") and len(clis) == 1:
                datos["cliente"] = clis[0]
            elif not any(c.lower() == str(datos.get("cliente", "")).strip().lower() for c in clis):
                raise HTTPException(status_code=403, detail="No tiene permisos para asignar este cliente")

    val = database.validar_inspeccion_candados(datos, excluir_id=inspeccion_id)
    if not val.get('valido', True):
        raise HTTPException(status_code=400, detail=val.get('mensaje', 'Violación de regla de seguridad'))

    ok = database.actualizar_inspeccion(inspeccion_id, datos)
    if not ok:
        raise HTTPException(status_code=404, detail="No se encontró el registro para actualizar")
    return {"status": "ok", "id": inspeccion_id, "message": "Inspección actualizada correctamente"}

@app.delete("/api/inspecciones/{inspeccion_id}")
def api_eliminar_inspeccion(inspeccion_id: int, usuario_actual: dict = Depends(obtener_usuario_actual)):
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    if es_rol_gerencia(rol):
        raise HTTPException(status_code=403, detail="El rol Gerencia tiene perfil ejecutivo de consulta y no puede eliminar inspecciones")
    if not es_rol_admin(rol, usuario_actual.get("username")):
        with database.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT cliente, usuario_registro FROM inspecciones WHERE id = ?", (inspeccion_id,))
            row_exist = cursor.fetchone()
            if not row_exist:
                raise HTTPException(status_code=404, detail="Inspección no encontrada")
            clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
            if clis and row_exist[0] and not any(c.lower() == str(row_exist[0]).strip().lower() for c in clis):
                raise HTTPException(status_code=403, detail="No tiene permisos para eliminar inspecciones de este cliente")

    database.eliminar_inspeccion(inspeccion_id)
    return {"status": "ok", "id": inspeccion_id, "message": "Inspección eliminada correctamente"}

# ==========================================

# RUTAS DE DIAGNÓSTICO Y PLANES DE ACCIÓN
# ==========================================

@app.get("/api/diagnostico")
def api_diagnostico(usuario: str = "", usuario_actual: dict = Depends(obtener_usuario_actual)):
    try:
        filtro_u = resolver_filtro_usuario_inspecciones(usuario_actual, usuario)
        df = database.obtener_todas_inspecciones(usuario=filtro_u)
        analisis = database.analizar_inspecciones_df(df)
        analisis = analisis.fillna("")
        return sanitizar_nativos(analisis.to_dict(orient="records"))
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Error al calcular diagnóstico: {str(e)}")


@app.get("/api/diagnostico/matriz-semaforos")
def api_diagnostico_matriz_semaforos(
    cliente: str = "",
    subcliente: str = "",
    operacion: str = "",
    placa: str = "",
    usuario: str = "",
    usuario_actual: dict = Depends(obtener_usuario_actual)
):
    try:
        rol = str(usuario_actual.get("rol", "")).strip().lower()
        user = str(usuario_actual.get("username", "")).strip().lower()
        if not es_rol_admin(rol, user):
            clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
            if clis and cliente and cliente != "Todos":
                if not any(c.lower() == cliente.strip().lower() for c in clis):
                    raise HTTPException(status_code=403, detail="No tiene permisos para ver la matriz de este cliente")

        filtro_u = resolver_filtro_usuario_inspecciones(usuario_actual, usuario)
        df = database.obtener_todas_inspecciones(usuario=filtro_u)
        if df.empty:
            return {
                "posiciones": [],
                "filas": [],
                "totales_posicion": {},
                "resumen_global": {"total_vehiculos": 0, "total_llantas": 0, "rojo": 0, "amarillo": 0, "verde": 0, "gris": 0}
            }

        if cliente and cliente != "Todos":
            df = df[df["cliente"].astype(str).str.upper().str.strip() == str(cliente).upper().strip()]
        if subcliente and subcliente != "Todos":
            df = df[df["subcliente"].astype(str).str.upper().str.strip() == str(subcliente).upper().strip()]
        if operacion and operacion != "Todos":
            df = df[df["tipo_operacion"].astype(str).str.upper().str.strip() == str(operacion).upper().strip()]
        if placa:
            df = df[df["placa"].astype(str).str.upper().str.contains(str(placa).upper().strip(), na=False)]

        matriz = database.calcular_matriz_semaforos_posicion(df)
        return sanitizar_nativos(matriz)
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Error al calcular matriz de semáforos: {str(e)}")



@app.get("/api/diagnostico/graficos")
def api_diagnostico_graficos(
    cliente: str = "",
    subcliente: str = "",
    tipo_operacion: str = "",
    tipologia: str = "",
    marca: str = "",
    dimension: str = "",
    eje_temporal: str = "mes",  # "mes" o "insp"
    usuario_actual: dict = Depends(obtener_usuario_actual)
):
    try:
        rol = str(usuario_actual.get("rol", "")).strip().lower()
        user = str(usuario_actual.get("username", "")).strip().lower()
        if not es_rol_admin(rol, user) and usuario_actual.get("cliente_asignado"):
            clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
            if clis and (not cliente or cliente == "Todos"):
                cliente = clis[0]
            elif clis and cliente and cliente != "Todos":
                if not any(c.lower() == cliente.strip().lower() for c in clis):
                    raise HTTPException(status_code=403, detail="No tiene permisos para ver los gráficos de este cliente")

        filtros = {
            "cliente": "" if cliente == "Todos" else cliente,
            "subcliente": "" if subcliente == "Todos" else subcliente,
            "tipo_operacion": "" if tipo_operacion == "Todos" else tipo_operacion,
            "tipologia": "" if tipologia == "Todos" else tipologia,
            "marca": "" if marca == "Todos" else marca,
            "dimension": "" if dimension == "Todos" else dimension,
            "eje_temporal": eje_temporal,
            "usuario": usuario_actual["username"] if not es_rol_admin(rol, user) else ""
        }
        datos = database.obtener_datos_graficos_diagnostico(filtros)
        return sanitizar_nativos(datos)
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Error al calcular gráficos de diagnóstico: {str(e)}")


def _estilizar_hoja_excel(ws, color_header="0F172A"):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color=color_header, end_color=color_header, fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    body_font = Font(name='Segoe UI', size=9, color='1E293B')
    zebra_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

    alerta_roja_fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
    alerta_roja_font = Font(name='Segoe UI', size=9, bold=True, color='991B1B')

    alerta_ambar_fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
    alerta_ambar_font = Font(name='Segoe UI', size=9, bold=True, color='92400E')

    ws.freeze_panes = 'A2'
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 28

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        for row_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 20
            is_even = (row_idx % 2 == 0)
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                val = cell.value
                cell.border = thin_border
                cell.font = body_font

                if is_even:
                    cell.fill = zebra_fill

                val_str = str(val or '').strip().upper()
                if val_str == 'SÍ':
                    col_name = str(ws.cell(row=1, column=col_idx).value or '').upper()
                    if 'CAMBIO' in col_name:
                        cell.fill = alerta_roja_fill
                        cell.font = alerta_roja_font
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    elif any(k in col_name for k in ['ALINEACION', 'ALINEACIÓN', 'DUAL']):
                        cell.fill = alerta_ambar_fill
                        cell.font = alerta_ambar_font
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                elif val_str in ('NO', 'OK'):
                    cell.alignment = Alignment(horizontal='center', vertical='center')

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                v = str(cell.value or '')
                if '\n' in v:
                    v = max(v.split('\n'), key=len)
                max_len = max(max_len, len(v))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 50)


@app.get("/api/diagnostico/exportar-excel")
def api_exportar_diagnostico_excel(
    cliente: str = "",
    subcliente: str = "",
    tipo_operacion: str = "",
    alerta: str = "",
    placa: str = "",
    usuario: str = "",
    usuario_actual: dict = Depends(obtener_usuario_actual)
):
    filtro_u = resolver_filtro_usuario_inspecciones(usuario_actual, usuario)
    df_insp = database.obtener_todas_inspecciones(usuario=filtro_u)

    # Validar permisos de asignación de cliente
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    if not es_rol_admin(rol, usuario_actual.get("username")):
        clis_permitidos = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
        if clis_permitidos and not df_insp.empty:
            df_insp = df_insp[df_insp['cliente'].astype(str).str.strip().str.lower().isin([c.lower() for c in clis_permitidos])]

    # 1. Filtros sobre inspecciones
    cli_limpio = str(cliente or '').strip()
    if cli_limpio and cli_limpio.upper() != 'TODOS' and not df_insp.empty:
        df_insp = df_insp[df_insp['cliente'].astype(str).str.upper().str.strip() == cli_limpio.upper()]

    sede_limpia = str(subcliente or '').strip()
    if sede_limpia and sede_limpia.upper() != 'TODOS' and not df_insp.empty:
        df_insp = df_insp[df_insp['subcliente'].astype(str).str.upper().str.strip() == sede_limpia.upper()]

    op_limpia = str(tipo_operacion or '').strip()
    if op_limpia and op_limpia.upper() != 'TODOS' and not df_insp.empty:
        df_insp = df_insp[df_insp['tipo_operacion'].astype(str).str.upper().str.strip() == op_limpia.upper()]

    placa_limpia = str(placa or '').strip().upper()
    if placa_limpia and not df_insp.empty:
        df_insp = df_insp[df_insp['placa'].astype(str).str.upper().str.contains(placa_limpia, na=False)]

    # 2. Análisis de diagnóstico
    df_analisis = database.analizar_inspecciones_df(df_insp)

    # 3. Filtro por tipo de alerta sobre df_analisis
    alerta_limpia = str(alerta or '').strip()
    if alerta_limpia and alerta_limpia.upper() != 'TODOS' and not df_analisis.empty:
        al_u = alerta_limpia.upper()
        if al_u in ('CAMBIO', 'SUGERENCIA CAMBIO'):
            df_analisis = df_analisis[df_analisis['Plan_Cambio_Llanta'] == 'SÍ']
        elif al_u in ('ALINEACION', 'ALINEACIÓN', 'SOLO PLAN ALINEACIÓN'):
            df_analisis = df_analisis[df_analisis['Plan_Alineacion'] == 'SÍ']
        elif al_u in ('DUAL', 'SOLO DUAL MAL APAREADO'):
            df_analisis = df_analisis[df_analisis['Plan_Dual_Mal_Apareado'] == 'SÍ']
        elif al_u in ('ALERTA', 'CUALQUIER ALERTA'):
            df_analisis = df_analisis[
                (df_analisis['Plan_Cambio_Llanta'] == 'SÍ') |
                (df_analisis['Plan_Alineacion'] == 'SÍ') |
                (df_analisis['Plan_Dual_Mal_Apareado'] == 'SÍ')
            ]
        elif al_u in ('OK', 'SIN NOVEDAD', 'SIN NOVEDADES', 'SIN NOVEDAD (OK)'):
            df_analisis = df_analisis[
                (df_analisis['Plan_Cambio_Llanta'] != 'SÍ') &
                (df_analisis['Plan_Alineacion'] != 'SÍ') &
                (df_analisis['Plan_Dual_Mal_Apareado'] != 'SÍ')
            ]

    # Sincronizar df_insp con los vehículos resultantes en df_analisis
    if not df_analisis.empty and not df_insp.empty:
        placas_filtradas = set(df_analisis['Vehículo/Placa'].dropna().astype(str).str.upper().unique())
        df_insp = df_insp[df_insp['placa'].astype(str).str.upper().isin(placas_filtradas)]
    elif df_analisis.empty and not df_insp.empty and alerta_limpia and alerta_limpia.upper() != 'TODOS':
        df_insp = df_insp.iloc[0:0]

    # Resumen de filtros aplicados para la hoja de auditoría
    df_resumen = pd.DataFrame([
        {"Parámetro": "Fecha de Generación", "Valor": ahora_bogota().strftime("%d/%m/%Y %H:%M:%S")},
        {"Parámetro": "Usuario Solicitante", "Valor": usuario_actual.get("username", "sistema")},
        {"Parámetro": "Filtro Cliente", "Valor": cli_limpio or "Todos los Clientes"},
        {"Parámetro": "Filtro Sede / Subcliente", "Valor": sede_limpia or "Todas las Sedes"},
        {"Parámetro": "Filtro Operación", "Valor": op_limpia or "Todas las Operaciones"},
        {"Parámetro": "Filtro Tipo Alerta", "Valor": alerta_limpia or "Todas las Alertas"},
        {"Parámetro": "Filtro Búsqueda Placa", "Valor": placa_limpia or "Sin filtro de placa"},
        {"Parámetro": "Total Vehículos en Diagnóstico", "Valor": len(df_analisis)},
        {"Parámetro": "Total Registros de Inspección", "Valor": len(df_insp)}
    ])

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df_analisis.to_excel(writer, sheet_name="Planes_Accion", index=False)
        if not df_insp.empty:
            df_insp.to_excel(writer, sheet_name="Detalle_Inspecciones", index=False)
        df_resumen.to_excel(writer, sheet_name="Filtros_Aplicados", index=False)

        wb = writer.book
        if "Planes_Accion" in wb.sheetnames:
            _estilizar_hoja_excel(wb["Planes_Accion"], color_header="0F172A")
        if "Detalle_Inspecciones" in wb.sheetnames:
            _estilizar_hoja_excel(wb["Detalle_Inspecciones"], color_header="1E293B")
        if "Filtros_Aplicados" in wb.sheetnames:
            _estilizar_hoja_excel(wb["Filtros_Aplicados"], color_header="047857")

    f_str = ahora_bogota().strftime("%Y%m%d_%H%M")
    tag_cli = f"_{cli_limpio.replace(' ', '_')}" if (cli_limpio and cli_limpio.upper() != 'TODOS') else ""
    tag_alerta = f"_{alerta_limpia.replace(' ', '_')}" if (alerta_limpia and alerta_limpia.upper() != 'TODOS') else ""
    filename = f"Diagnostico{tag_cli}{tag_alerta}_{f_str}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# ==========================================
# RUTAS DE TRAZABILIDAD Y KPIS DE LLANTAS
# ==========================================

@app.get("/api/kpis-llantas")
def api_kpis_llantas(
    numero_interno: str = "",
    cliente: str = "",
    subcliente: str = "",
    tipo_operacion: str = "",
    marca: str = "",
    diseno: str = "",
    dimension: str = "",
    vida: str = "",
    banda_reen: str = "",
    usuario: str = "",
    usuario_actual: dict = Depends(obtener_usuario_actual)
):
    filtros = {
        'cliente': cliente,
        'subcliente': subcliente,
        'tipo_operacion': tipo_operacion,
        'marca': marca,
        'diseno': diseno,
        'dimension': dimension,
        'vida': vida,
        'banda_reen': banda_reen,
        'usuario': usuario
    }
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    user = str(usuario_actual.get("username", "")).strip().lower()
    es_admin = es_rol_admin(rol, user)
    es_privilegiado = es_admin or es_rol_jefatura(rol) or es_rol_gerencia(rol)
    if not es_privilegiado:
        filtros['usuario'] = usuario_actual["username"]
    elif usuario:
        filtros['usuario'] = usuario
    elif usuario_actual.get("cliente_asignado"):
        filtros['usuario'] = usuario_actual["username"]

    raw = database.obtener_trazabilidad_kpis_llanta(numero_interno, filtros)
    return sanitizar_nativos(raw)

@app.get("/api/kpis-llantas/exportar-excel")
def api_exportar_kpis_excel(
    numero_interno: str = "",
    cliente: str = "",
    subcliente: str = "",
    tipo_operacion: str = "",
    marca: str = "",
    diseno: str = "",
    dimension: str = "",
    vida: str = "",
    banda_reen: str = "",
    usuario: str = "",
    usuario_actual: dict = Depends(obtener_usuario_actual)
):
    filtros = {
        'cliente': cliente,
        'subcliente': subcliente,
        'tipo_operacion': tipo_operacion,
        'marca': marca,
        'diseno': diseno,
        'dimension': dimension,
        'vida': vida,
        'banda_reen': banda_reen,
        'usuario': usuario
    }
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    user = str(usuario_actual.get("username", "")).strip().lower()
    es_admin = es_rol_admin(rol, user)
    es_privilegiado = es_admin or es_rol_jefatura(rol) or es_rol_gerencia(rol)
    if not es_privilegiado:
        filtros['usuario'] = usuario_actual["username"]
    elif usuario:
        filtros['usuario'] = usuario
    elif usuario_actual.get("cliente_asignado"):
        filtros['usuario'] = usuario_actual["username"]

    data = database.obtener_trazabilidad_kpis_llanta(numero_interno, filtros)
    df = pd.DataFrame(data)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Trazabilidad_KPIs", index=False)
    f_str = ahora_bogota().strftime("%Y%m%d_%H%M")
    filename = f"KPIs_Llantas_TireAnalyze_{f_str}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# ==========================================
# RUTAS DE REPORTES PDF
# ==========================================

@app.get("/api/informe-pdf/{placa}")
def api_descargar_pdf_plan_mantenimiento(placa: str, fecha: str = "", usuario_actual: dict = Depends(obtener_usuario_actual)):
    try:
        pdf_bytes = pdf_generator.generar_pdf_plan_mantenimiento(placa, fecha if fecha else None)
        f_tag = fecha.replace("-", "") if fecha else ahora_bogota().strftime('%Y%m%d')
        filename = f"TireAnalyze_Informe_{placa}_{f_tag}.pdf"
        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={"Content-Disposition": f"inline; filename={filename}"}
        )
    except Exception as e:
        raise HTTPException(status_code=404, detail=str(e))

@app.get("/api/informe-pdf-consolidado")
def api_descargar_pdf_consolidado(cliente: str = "", subcliente: str = "", fecha: str = "", operacion: str = "", usuario_actual: dict = Depends(obtener_usuario_actual)):
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    user = str(usuario_actual.get("username", "")).strip().lower()
    filtro_u = None
    if not es_rol_admin(rol, user):
        clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
        if clis:
            filtro_u = usuario_actual["username"]
            if cliente and cliente != "Todos":
                if not any(c.lower() == cliente.strip().lower() for c in clis):
                    raise HTTPException(status_code=403, detail="No tiene permisos para generar informe de este cliente")
            elif len(clis) == 1:
                cliente = clis[0]
    try:
        pdf_bytes = pdf_generator.generar_pdf_consolidado_jornada(
            cliente=cliente if (cliente and cliente != "Todos") else None,
            subcliente=subcliente if (subcliente and subcliente != "Todos") else None,
            fecha=fecha if (fecha and fecha != "Todos") else None,
            operacion=operacion if (operacion and operacion != "Todos") else None,
            usuario=filtro_u
        )
        c_tag = cliente.replace(" ", "_") if (cliente and cliente != "Todos") else "General"
        if subcliente and subcliente != "Todos":
            c_tag += f"_{subcliente.replace(' ', '_')}"
        if operacion and operacion != "Todos":
            c_tag += f"_{operacion.replace(' ', '_')}"
        f_tag = fecha.replace("-", "") if fecha else ahora_bogota().strftime('%Y%m%d')
        filename = f"TireAnalyze_Consolidado_Flota_{c_tag}_{f_tag}.pdf"
        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={"Content-Disposition": f"inline; filename={filename}"}
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=404, detail=str(e))

# ==========================================
# RUTAS DE CATÁLOGOS Y CONFIGURACIONES
# ==========================================

@app.get("/api/posiciones")
def api_posiciones(tipologia: str, usuario_actual: dict = Depends(obtener_usuario_actual)):
    posiciones = database.obtener_posiciones_por_tipologia(tipologia)
    return {"tipologia": tipologia, "posiciones": posiciones}

@app.get("/api/operaciones")
def api_obtener_operaciones(usuario_actual: dict = Depends(obtener_usuario_actual)):
    return database.obtener_todas_operaciones()

@app.post("/api/operaciones")
def api_crear_operacion(item: NuevaOperacionItem, admin: dict = Depends(requerir_admin_o_jefatura)):
    nom = database.guardar_operacion(item.nombre)
    return {"status": "ok", "nombre": nom}

@app.put("/api/operaciones")
def api_editar_operacion(item: EditarOperacionItem, admin: dict = Depends(requerir_admin_o_jefatura)):
    ok = database.actualizar_operacion(item.nombre_antiguo, item.nombre_nuevo)
    return {"status": "ok" if ok else "error"}

@app.delete("/api/operaciones")
def api_eliminar_operacion(nombre: str, admin: dict = Depends(requerir_admin_o_jefatura)):
    ok = database.eliminar_operacion(nombre)
    return {"status": "ok" if ok else "error"}

@app.get("/api/tipologias")
def api_obtener_tipologias(usuario_actual: dict = Depends(obtener_usuario_actual)):
    return database.obtener_todas_tipologias()

@app.post("/api/tipologias")
def api_crear_tipologia(item: NuevaTipologiaItem, admin: dict = Depends(requerir_admin_o_jefatura)):
    nom = database.guardar_tipologia(
        nombre=item.nombre,
        descripcion=item.descripcion,
        cant_llantas=item.cant_llantas,
        esquema_custom=item.esquema_custom,
        ejes_dir=item.ejes_dir,
        tipo_dir=item.tipo_dir,
        ejes_trac=item.ejes_trac,
        tipo_trac=item.tipo_trac,
        ejes_libres=item.ejes_libres,
        tipo_libres=item.tipo_libres,
        cant_repuestos=item.cant_repuestos
    )
    return {"status": "ok", "nombre": nom}

@app.put("/api/tipologias")
def api_editar_tipologia(item: EditarTipologiaItem, admin: dict = Depends(requerir_admin_o_jefatura)):
    try:
        ok = database.actualizar_tipologia(
            nombre_antiguo=item.nombre_antiguo,
            nombre_nuevo=item.nombre_nuevo,
            descripcion=item.descripcion or "",
            cant_llantas=item.cant_llantas,
            ejes_dir=item.ejes_dir,
            tipo_dir=item.tipo_dir or "sencilla",
            ejes_trac=item.ejes_trac,
            tipo_trac=item.tipo_trac or "dual",
            ejes_libres=item.ejes_libres,
            tipo_libres=item.tipo_libres or "dual",
            cant_repuestos=item.cant_repuestos,
            esquema_custom=item.esquema_custom
        )
        return {"status": "ok" if ok else "error"}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.delete("/api/tipologias")
def api_eliminar_tipologia(nombre: str, admin: dict = Depends(requerir_admin_o_jefatura)):
    try:
        ok = database.eliminar_tipologia(nombre)
        return {"status": "ok" if ok else "error"}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.get("/api/esquema")
def api_esquema(tipologia: str, usuario_actual: dict = Depends(obtener_usuario_actual)):
    esquema = database.obtener_esquema_tipologia(tipologia)
    posiciones = database.obtener_posiciones_por_tipologia(tipologia)
    return {"tipologia": tipologia, "esquema": esquema, "posiciones": posiciones}

@app.get("/api/modelos-llanta")
def api_modelos_llanta(usuario_actual: dict = Depends(obtener_usuario_actual)):
    return database.obtener_catalogo_modelos()

@app.post("/api/modelos-llanta")
def api_guardar_modelo(item: ModeloLlantaItem, admin: dict = Depends(requerir_admin_o_jefatura)):
    ok = database.guardar_modelo_llanta(item.tipo_llanta, item.marca, item.diseno, item.dimension, item.prof_inicial, item.descripcion)
    return {"status": "ok" if ok else "error"}

@app.put("/api/modelos-llanta")
def api_editar_modelo_llanta(item: EditarModeloLlantaItem, admin: dict = Depends(requerir_admin_o_jefatura)):
    try:
        res = database.actualizar_modelo_llanta(
            modelo_id=item.id,
            tipo_llanta=item.tipo_llanta,
            marca=item.marca,
            diseno=item.diseno,
            dimension=item.dimension,
            prof_inicial=item.prof_inicial,
            descripcion=item.descripcion or "",
            actualizar_bd=item.actualizar_bd if item.actualizar_bd is not None else True
        )
        return res
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.delete("/api/modelos-llanta/{modelo_id}")
def api_eliminar_modelo_llanta(modelo_id: int, admin: dict = Depends(requerir_admin_o_jefatura)):
    try:
        ok = database.eliminar_modelo_llanta(modelo_id)
        return {"status": "ok" if ok else "error"}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.get("/api/profundidad-inicial")
def api_buscar_profundidad(marca: str = "", diseno: str = "", dimension: str = "", vida: str = "Original", banda: str = "", usuario_actual: dict = Depends(obtener_usuario_actual)):
    prof = database.buscar_profundidad_inicial(marca, diseno, dimension, vida, banda=banda)
    return {"prof_inicial": prof}

@app.get("/api/buscar-modelos")
def api_buscar_modelos(query: str = "", usuario_actual: dict = Depends(obtener_usuario_actual)):
    return database.buscar_modelos_llantas_filtro(query)

@app.get("/api/clientes-directorio")
def api_clientes_directorio(usuario_actual: dict = Depends(obtener_usuario_actual)):
    directorio = database.obtener_directorio_clientes()
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    user = str(usuario_actual.get("username", "")).strip().lower()
    if not es_rol_admin(rol, user) and usuario_actual.get("cliente_asignado"):
        clis = database.parsear_clientes_asignados(usuario_actual["cliente_asignado"])
        directorio = [c for c in directorio if c["cliente"] in clis]
    return sanitizar_nativos(directorio)

@app.post("/api/clientes-directorio")
def api_crear_cliente_directorio(item: ClienteDirectorioItem, admin: dict = Depends(requerir_admin_o_jefatura)):
    try:
        res = database.crear_o_actualizar_cliente(item.model_dump())
        if not res.get("ok"):
            raise HTTPException(status_code=400, detail=res.get("mensaje", "Error al procesar el cliente"))
        return sanitizar_nativos({
            "status": "ok",
            "ok": True,
            "cliente": res.get("cliente"),
            "mensaje": res.get("mensaje")
        })
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al guardar cliente: {str(e)}")

@app.delete("/api/clientes-directorio/{cliente}")
def api_eliminar_cliente_directorio(cliente: str, admin: dict = Depends(requerir_admin_o_jefatura)):
    res = database.eliminar_cliente_directorio(cliente)
    if not res.get("ok"):
        raise HTTPException(status_code=400, detail=res.get("mensaje", "No se pudo eliminar el cliente"))
    return sanitizar_nativos({
        "status": "ok",
        "ok": True,
        "mensaje": res.get("mensaje")
    })

@app.get("/api/configuraciones-clientes")
def api_obtener_configuraciones_clientes(admin: dict = Depends(requerir_admin_o_jefatura)):
    return database.obtener_todas_configuraciones_clientes()

@app.get("/api/limite-cliente")
def api_obtener_limite_cliente(cliente: str = "", subcliente: str = "TODOS", eje: str = "TODOS", usuario_actual: dict = Depends(obtener_usuario_actual)):
    lim = database.obtener_limite_retiro_cliente(cliente, subcliente, eje)
    lic = database.obtener_licencia_cliente(cliente)
    return {"cliente": cliente, "subcliente": subcliente, "eje": eje, "limite_retiro_mm": lim, "tipo_licencia": lic}

@app.post("/api/configuraciones-clientes")
def api_guardar_configuracion_cliente(item: LimiteClienteItem, admin: dict = Depends(requerir_admin_o_jefatura)):
    ok = database.guardar_limite_retiro_cliente(
        cliente=item.cliente,
        limite_retiro_mm=item.limite_retiro_mm,
        observaciones=item.observaciones,
        subcliente=item.subcliente,
        eje=item.eje
    )
    return {"status": "ok" if ok else "error"}

@app.put("/api/configuraciones-clientes")
def api_editar_configuracion_cliente(item: EditarLimiteClienteItem, admin: dict = Depends(requerir_admin_o_jefatura)):
    try:
        ok = database.actualizar_limite_retiro_cliente(
            limite_id=item.id,
            cliente=item.cliente,
            limite_retiro_mm=item.limite_retiro_mm,
            observaciones=item.observaciones or "",
            subcliente=item.subcliente or "TODOS",
            eje=item.eje or "TODOS"
        )
        return {"status": "ok" if ok else "error"}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.put("/api/configuraciones-clientes/{limite_id}")
def api_editar_configuracion_cliente_por_id(limite_id: int, item: LimiteClienteItem, admin: dict = Depends(requerir_admin_o_jefatura)):
    try:
        ok = database.actualizar_limite_retiro_cliente(
            limite_id=limite_id,
            cliente=item.cliente,
            limite_retiro_mm=item.limite_retiro_mm,
            observaciones=item.observaciones or "",
            subcliente=item.subcliente or "TODOS",
            eje=item.eje or "TODOS"
        )
        return {"status": "ok" if ok else "error"}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.delete("/api/configuraciones-clientes/{limite_id}")
def api_eliminar_configuracion_cliente(limite_id: int, admin: dict = Depends(requerir_admin_o_jefatura)):
    ok = database.eliminar_limite_retiro_cliente(limite_id)
    return sanitizar_nativos({"status": "ok" if ok else "error"})

# ==========================================
# RUTAS DE MARCAS POR CLIENTE
# ==========================================

@app.get("/api/configuracion-marcas-cliente")
def api_obtener_marcas_cliente(cliente: str = "", usuario_actual: dict = Depends(obtener_usuario_actual)):
    if usuario_actual.get("rol") != "admin":
        clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
        if clis:
            if cliente and not any(c.lower() == cliente.strip().lower() for c in clis):
                raise HTTPException(status_code=403, detail="No tiene permisos para ver configuraciones de este cliente")
            if not cliente and len(clis) == 1:
                cliente = clis[0]
    if cliente:
        res = database.obtener_marcas_por_cliente(cliente)
        return sanitizar_nativos({"cliente": cliente.strip().upper(), "marcas": res})
    todas = database.obtener_todas_configuraciones_marcas()
    if usuario_actual.get("rol") != "admin":
        clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
        if clis:
            todas = [t for t in todas if any(c.lower() == t["cliente"].strip().lower() for c in clis)]
    return sanitizar_nativos(todas)

@app.post("/api/configuracion-marcas-cliente")
def api_guardar_marca_cliente(item: MarcaClienteItem, usuario_actual: dict = Depends(obtener_usuario_actual)):
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    cli = str(item.cliente or "").strip().upper()
    m = str(item.marca or "").strip()
    if not cli or not m:
        raise HTTPException(status_code=400, detail="Cliente y marca son obligatorios")
    if rol not in ("admin", "jefatura"):
        clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
        if clis and not any(c.lower() == cli.lower() for c in clis):
            raise HTTPException(status_code=403, detail="No tiene permisos para parametrizar marcas en este cliente")
    database.guardar_marca_cliente(cli, m)
    return sanitizar_nativos({"status": "ok", "mensaje": f"Marca '{m}' asociada a '{cli}' exitosamente"})

@app.delete("/api/configuracion-marcas-cliente/{marca_id}")
def api_eliminar_marca_cliente_id(marca_id: int, admin: dict = Depends(requerir_admin_o_jefatura)):
    ok = database.eliminar_marca_cliente_por_id(marca_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Configuración no encontrada")
    return sanitizar_nativos({"status": "ok", "mensaje": "Marca eliminada del cliente"})

@app.delete("/api/configuracion-marcas-cliente")
def api_eliminar_marca_cliente(cliente: str, marca: str, admin: dict = Depends(requerir_admin_o_jefatura)):
    ok = database.eliminar_marca_cliente(cliente, marca)
    if not ok:
        raise HTTPException(status_code=404, detail="Configuración no encontrada")
    return sanitizar_nativos({"status": "ok", "mensaje": "Marca desvinculada del cliente"})

@app.post("/api/configuraciones-clientes/licencia")
def api_guardar_licencia_cliente(item: ConfigLicenciaClienteItem, admin: dict = Depends(requerir_admin)):
    ok = database.guardar_licencia_cliente(item.cliente, item.tipo_licencia, item.plantilla_kpis_json)
    if not ok:
        raise HTTPException(status_code=400, detail="Error al guardar la configuración de licencia")
    return sanitizar_nativos({"status": "ok"})

@app.post("/api/admin/migrar-pepsico-tellantas")
@app.get("/api/admin/migrar-pepsico-tellantas")
def api_migrar_pepsico_tellantas(usuario_actual: dict = Depends(requerir_admin_o_jefatura)):
    res = database.migrar_pepsico_a_tellantas()
    return sanitizar_nativos(res)

@app.get("/api/valores-campo")
def api_valores_campo(campo: str, cliente: str = "", usuario: str = "", marca: str = "", todas: bool = False, usuario_actual: dict = Depends(obtener_usuario_actual)):
    if usuario_actual.get("rol") != "admin":
        clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
        if campo == "cliente" and clis:
            return sanitizar_nativos(sorted(list(set(clis))))
        if cliente and clis and not any(c.lower() == cliente.strip().lower() for c in clis):
            return sanitizar_nativos([])
        if not cliente and len(clis) == 1:
            cliente = clis[0]
    elif usuario and not cliente:
        usr_info = database.obtener_usuario_por_username(usuario)
        if usr_info and usr_info.get("rol") != "admin":
            clis = database.parsear_clientes_asignados(usr_info.get("cliente_asignado"))
            if campo == "cliente" and clis:
                return sanitizar_nativos(sorted(list(set(clis))))
            if cliente and clis and not any(c.lower() == cliente.strip().lower() for c in clis):
                return sanitizar_nativos([])
            if not cliente and len(clis) == 1:
                cliente = clis[0]
    res = database.obtener_valores_unicos_campo(campo, cliente, filtro_marca=marca, todas=todas)
    return sanitizar_nativos(res)

@app.get("/api/vehiculo-perfil")
def api_vehiculo_perfil(placa: str, usuario_actual: dict = Depends(obtener_usuario_actual)):
    perfil = database.obtener_perfil_vehiculo(placa)
    # Validar permisos de inspector con clientes asignados
    if usuario_actual.get("rol") != "admin":
        clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
        if clis and perfil.get("encontrado") and perfil.get("cliente"):
            if not any(c.lower() == perfil["cliente"].strip().lower() for c in clis):
                raise HTTPException(status_code=403, detail="No tiene permisos para ver vehículos de este cliente")
    return sanitizar_nativos(perfil)

@app.get("/api/llanta-por-cliente")
def api_llanta_por_cliente(
    numero_interno: str,
    cliente: str = "",
    placa: str = "",
    usuario_actual: dict = Depends(obtener_usuario_actual)
):
    if usuario_actual.get("rol") != "admin":
        clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
        if clis:
            if cliente and not any(c.lower() == cliente.strip().lower() for c in clis):
                raise HTTPException(status_code=403, detail="No tiene permisos para consultar llantas de este cliente")
            if not cliente and len(clis) == 1:
                cliente = clis[0]
    res = database.obtener_llanta_por_cliente(cliente=cliente, numero_interno=numero_interno, placa=placa)
    if res:
        return sanitizar_nativos({"encontrado": True, "datos": res})
    return {"encontrado": False}

@app.get("/api/inspeccion-anterior")
def api_inspeccion_anterior(placa: str, posicion: int, usuario_actual: dict = Depends(obtener_usuario_actual)):
    prev = database.obtener_ultima_inspeccion_posicion(placa, posicion)
    perfil = database.obtener_perfil_vehiculo(placa)
    if prev:
        return sanitizar_nativos({
            "encontrado": True,
            "datos": prev,
            "vehiculo_perfil": perfil if perfil.get("encontrado") else None
        })
    return {"encontrado": False, "vehiculo_perfil": perfil if perfil.get("encontrado") else None}

@app.get("/api/posiciones-inspeccionadas")
def api_posiciones_inspeccionadas(placa: str, fecha: str = "", usuario_actual: dict = Depends(obtener_usuario_actual)):
    det = database.obtener_posiciones_inspeccionadas_detalle(placa, fecha)
    return {
        "placa": placa,
        "posiciones": det["posiciones"],
        "posiciones_con_novedad": det["posiciones_con_novedad"],
        "novedades": det["novedades"]
    }

# ==========================================
# RUTAS DE AUTENTICACIÓN Y USUARIOS
# ==========================================

@app.post("/api/login")
def api_login(item: LoginItem, response: Response):
    usr = database.autenticar_usuario(item.username, item.password)
    if not usr:
        raise HTTPException(status_code=401, detail="Usuario o contraseña incorrectos")
    token = crear_token_acceso(usr["username"], usr["rol"])
    response.set_cookie(
        key="tireanalyze_token",
        value=token,
        httponly=True,
        samesite="lax",
        max_age=TOKEN_EXPIRATION_SECONDS
    )
    response.set_cookie(
        key="tirecheck_token",
        value=token,
        httponly=True,
        samesite="lax",
        max_age=TOKEN_EXPIRATION_SECONDS
    )
    return {"status": "ok", "token": token, "usuario": usr}

@app.get("/api/usuario-actual")
def api_usuario_actual(usuario_actual: dict = Depends(obtener_usuario_actual)):
    return sanitizar_nativos({
        "status": "ok",
        "usuario": usuario_actual
    })

@app.post("/api/logout")
def api_logout(response: Response):
    response.delete_cookie(key="tireanalyze_token")
    response.delete_cookie(key="tirecheck_token")
    return {"status": "ok", "mensaje": "Sesión cerrada correctamente"}

@app.get("/api/usuarios")
def api_listar_usuarios(admin: dict = Depends(requerir_admin)):
    return database.listar_usuarios()

@app.post("/api/usuarios")
def api_crear_usuario(item: UsuarioItem, admin: dict = Depends(requerir_admin)):
    ok = database.crear_usuario(
        username=item.username,
        password=item.password,
        nombre_completo=item.nombre_completo,
        rol=item.rol,
        cliente_asignado=item.cliente_asignado,
        vistas_permitidas=item.vistas_permitidas
    )
    if not ok:
        raise HTTPException(status_code=400, detail="El nombre de usuario ya existe o los datos son inválidos (contraseña mínima 4 caracteres)")
    return {"status": "ok", "username": item.username}

@app.put("/api/usuarios/{username}")
def api_actualizar_usuario(username: str, item: EditarUsuarioItem, admin: dict = Depends(requerir_admin)):
    if username.lower() == "admin":
        if item.activo is False:
            raise HTTPException(status_code=400, detail="No se puede desactivar al usuario administrador principal")
        if item.rol and item.rol.lower() != "admin":
            raise HTTPException(status_code=400, detail="No se puede degradar el rol del usuario administrador principal")
    ok = database.actualizar_usuario(
        username=username,
        nombre_completo=item.nombre_completo,
        password=item.password,
        rol=item.rol,
        cliente_asignado=item.cliente_asignado,
        activo=item.activo,
        vistas_permitidas=item.vistas_permitidas
    )
    if not ok:
        raise HTTPException(status_code=400, detail="No se pudo actualizar el usuario")
    return {"status": "ok", "username": username}

@app.post("/api/usuarios/{username}/reset-password")
def api_reset_password(username: str, item: ResetPasswordItem, admin: dict = Depends(requerir_admin)):
    if not item.nueva_password or len(item.nueva_password.strip()) < 4:
        raise HTTPException(status_code=400, detail="La nueva contraseña debe tener al menos 4 caracteres")
    usr_existente = database.obtener_usuario_por_username(username)
    if not usr_existente:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    ok = database.restablecer_password_usuario(username, item.nueva_password)
    if not ok:
        raise HTTPException(status_code=400, detail="No se pudo restablecer la contraseña")
    return {"status": "ok", "message": f"Contraseña del usuario {username} restablecida correctamente"}

@app.delete("/api/usuarios/{username}")
def api_eliminar_usuario(username: str, admin: dict = Depends(requerir_admin)):
    if username.lower() == "admin":
        raise HTTPException(status_code=400, detail="No se puede eliminar el usuario administrador principal")
    ok = database.eliminar_usuario(username)
    if not ok:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    return {"status": "ok", "username": username}

@app.get("/api/admin/db-info")
def api_admin_db_info(admin: dict = Depends(requerir_admin)):
    with database.get_connection() as conn:
        c = conn.cursor()
        total_insp = c.execute("SELECT COUNT(*) FROM inspecciones").fetchone()[0]
        total_usr = c.execute("SELECT COUNT(*) FROM usuarios").fetchone()[0]
    return {
        "db_path": database.DB_PATH,
        "db_size_bytes": os.path.getsize(database.DB_PATH) if os.path.exists(database.DB_PATH) else 0,
        "total_inspecciones": total_insp,
        "total_usuarios": total_usr
    }

@app.post("/api/admin/restaurar-seed")
def api_admin_restaurar_seed(admin: dict = Depends(requerir_admin)):
    import shutil
    candidates = [database._DB_SEED, database._DB_ROOT]
    seed_found = None
    for c in candidates:
        if os.path.exists(c) and os.path.getsize(c) > 10000:
            seed_found = c
            break
    if not seed_found:
        raise HTTPException(status_code=404, detail="No se encontró el archivo de semilla de la base de datos.")
    try:
        shutil.copy2(seed_found, database.DB_PATH)
        with database.get_connection() as conn:
            cur = conn.cursor()
            total_insp = cur.execute("SELECT COUNT(*) FROM inspecciones").fetchone()[0]
        return {
            "status": "ok",
            "message": f"Base de datos restaurada con éxito desde {seed_found}. Total inspecciones: {total_insp}",
            "total_inspecciones": total_insp
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al restaurar semilla: {str(e)}")

# ==========================================
# MÓDULO DE LAS 7 BODEGAS, KARDEX Y BALANCE MRP (KPI 20)
# ==========================================

class AltaLlantaBodegaItem(BaseModel):
    cliente: str
    subcliente: str = ""
    numero_interno: str
    ubicacion_bodega: str = "BODEGA_1_NUEVAS"
    marca: str
    diseno: str = ""
    dimension: str = "295/80R22.5"
    vida_actual: str = "Original"
    banda_reen: str = ""
    prof_inicial: float = 18.0
    prof_actual_min: float = 18.0
    costo_compra: float = 1450000.0
    proveedor_actual: str = ""
    observaciones: str = ""

class MovimientoBodegaItem(BaseModel):
    cliente: str
    numero_interno: str
    bodega_destino: str
    fecha_movimiento: str = ""
    placa: str = ""
    posicion: int = 0
    prof_min_mm: Optional[float] = None
    nueva_prof_inicial: Optional[float] = None
    banda_reen: str = ""
    km_vehiculo: float = 0.0
    costo_asociado: float = 0.0
    proveedor: str = ""
    causal_o_motivo: str = ""
    documento_soporte: str = ""

class GestionGarantiaItem(BaseModel):
    cliente: str
    numero_interno: str
    estado_garantia: str  # 'Posible Garantía', 'En Reclamación', 'Nota Crédito Reconocida', 'Rechazada'
    valor_reconocido_garantia: float = 0.0
    observaciones: str = ""

@app.get("/api/bodegas/resumen-mrp")
def api_bodegas_resumen_mrp(
    cliente: str = "",
    subcliente: str = "",
    dimension: str = "",
    usuario_actual: dict = Depends(obtener_usuario_actual)
):
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    user = str(usuario_actual.get("username", "")).strip().lower()
    if not es_rol_admin(rol, user) and usuario_actual.get("cliente_asignado"):
        clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
        if clis and (not cliente or cliente == "Todos"):
            cliente = clis[0]
        elif clis and cliente and cliente != "Todos":
            if not any(c.lower() == cliente.strip().lower() for c in clis):
                raise HTTPException(status_code=403, detail="No tiene permisos para ver las bodegas de este cliente")

    filtros = {
        "cliente": "" if cliente == "Todos" else cliente,
        "subcliente": "" if subcliente == "Todos" else subcliente,
        "dimension": "" if dimension == "Todos" else dimension
    }
    return sanitizar_nativos(database.obtener_balance_bodegas_mrp(filtros))

@app.get("/api/bodegas/inventario")
def api_bodegas_inventario(
    cliente: str = "",
    subcliente: str = "",
    ubicacion_bodega: str = "",
    dimension: str = "",
    q: str = "",
    usuario_actual: dict = Depends(obtener_usuario_actual)
):
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    user = str(usuario_actual.get("username", "")).strip().lower()
    if not es_rol_admin(rol, user) and usuario_actual.get("cliente_asignado"):
        clis = database.parsear_clientes_asignados(usuario_actual.get("cliente_asignado"))
        if clis and (not cliente or cliente == "Todos"):
            cliente = clis[0]
        elif clis and cliente and cliente != "Todos":
            if not any(c.lower() == cliente.strip().lower() for c in clis):
                raise HTTPException(status_code=403, detail="No tiene permisos para ver el inventario de este cliente")

    filtros = {
        "cliente": "" if cliente == "Todos" else cliente,
        "subcliente": "" if subcliente == "Todos" else subcliente,
        "ubicacion_bodega": "" if ubicacion_bodega == "Todas" else ubicacion_bodega,
        "dimension": "" if dimension == "Todas" else dimension,
        "q": q
    }
    return sanitizar_nativos(database.listar_inventario_bodegas(filtros))

@app.post("/api/bodegas/alta")
def api_bodegas_alta(item: AltaLlantaBodegaItem, usuario_actual: dict = Depends(obtener_usuario_actual)):
    datos = item.model_dump()
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    user = str(usuario_actual.get("username", "")).strip().lower()
    if not es_rol_admin(rol, user) and usuario_actual.get("cliente_asignado"):
        datos["cliente"] = usuario_actual["cliente_asignado"]
    datos["usuario_registro"] = usuario_actual.get("username", "admin")
    res = database.crear_llanta_en_bodega(datos)
    if not res.get("ok"):
        raise HTTPException(status_code=400, detail=res.get("mensaje", "Error al dar de alta la llanta"))
    return sanitizar_nativos(res)

@app.post("/api/bodegas/movimiento")
def api_bodegas_movimiento(item: MovimientoBodegaItem, usuario_actual: dict = Depends(obtener_usuario_actual)):
    datos = item.model_dump()
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    user = str(usuario_actual.get("username", "")).strip().lower()
    if not es_rol_admin(rol, user) and usuario_actual.get("cliente_asignado"):
        datos["cliente"] = usuario_actual["cliente_asignado"]
    datos["usuario_registro"] = usuario_actual.get("username", "admin")
    res = database.registrar_movimiento_bodega(datos)
    if not res.get("ok"):
        raise HTTPException(status_code=400, detail=res.get("mensaje", "Movimiento de bodega inválido"))
    return sanitizar_nativos(res)

@app.post("/api/bodegas/garantia")
def api_bodegas_garantia(item: GestionGarantiaItem, usuario_actual: dict = Depends(obtener_usuario_actual)):
    datos = item.model_dump()
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    user = str(usuario_actual.get("username", "")).strip().lower()
    if not es_rol_admin(rol, user) and usuario_actual.get("cliente_asignado"):
        datos["cliente"] = usuario_actual["cliente_asignado"]
    datos["usuario_registro"] = usuario_actual.get("username", "admin")
    res = database.actualizar_garantia_llanta(datos)
    if not res.get("ok"):
        raise HTTPException(status_code=400, detail=res.get("mensaje", "No se pudo actualizar la garantía"))
    return sanitizar_nativos(res)

@app.get("/api/bodegas/kardex")
def api_bodegas_kardex(cliente: str = "", numero_interno: str = "", usuario_actual: dict = Depends(obtener_usuario_actual)):
    rol = str(usuario_actual.get("rol", "")).strip().lower()
    user = str(usuario_actual.get("username", "")).strip().lower()
    if not es_rol_admin(rol, user) and usuario_actual.get("cliente_asignado"):
        cliente = usuario_actual["cliente_asignado"]
    return sanitizar_nativos(database.obtener_kardex_llanta(cliente, numero_interno))

# ==========================================
# PUNTO DE ENTRADA
# ==========================================

if __name__ == "__main__":
    import uvicorn
    raw_port = os.environ.get("PORT", "8000")
    try:
        puerto = int(raw_port)
    except (ValueError, TypeError):
        puerto = 8000
    recarga = os.environ.get("ENV") != "production"
    uvicorn.run("server:app", host="0.0.0.0", port=puerto, reload=recarga)
